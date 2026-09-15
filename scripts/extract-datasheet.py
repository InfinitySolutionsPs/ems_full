#!/usr/bin/env python3
import json, sys
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook

source = Path(sys.argv[1])
target = Path(sys.argv[2])
sheet = load_workbook(source, data_only=True, read_only=True)["Data Sheet"]

def txt(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def num(v):
    try: return float(v) if v not in (None, "") else None
    except: return None

def day(v):
    if not isinstance(v, (date, datetime)): return ""
    # The workbook mixes dd/mm and mm/dd serials. Its reporting period is 1–8 Sep 2026.
    if v.year == 2026 and v.month != 9 and v.day == 9:
        return f"2026-09-{v.month:02d}"
    return v.strftime("%Y-%m-%d")

def clock(v):
    if isinstance(v, (datetime, time)): return v.strftime("%H:%M")
    if isinstance(v, (int, float)):
        mins = round((v % 1) * 1440) % 1440
        return f"{mins//60:02d}:{mins%60:02d}"
    s=txt(v)
    return s[:5] if s else ""

def yes(v): return txt(v).lower() in {"yes", "y", "true", "1", "نعم", "x", "✓"}

raw=[]
for r, cells in enumerate(sheet.iter_rows(min_row=3, max_col=64), start=3):
    vals=[cell.value for cell in cells]
    if not any(v not in (None, "") for v in vals[1:57]): continue
    if not (txt(vals[8]) and day(vals[23])): continue
    raw.append((r,vals))

numbers=Counter(txt(v[1]) for _,v in raw if txt(v[1]))
codes=Counter(txt(v[2]) for _,v in raw if txt(v[2]))
records=[]
for row,v in raw:
    original_number, original_code = txt(v[1]), txt(v[2])
    incident_number = original_number if original_number not in {"", "0"} and numbers[original_number] == 1 else f"XLS-DS-{row:04d}"
    incident_code = original_code if original_code and codes[original_code] == 1 else f"XLS-DS-{row:04d}"
    start,end,total=num(v[20]),num(v[21]),num(v[22])
    distance=total if total is not None and total >= 0 else ((end-start) if start is not None and end is not None and end >= start else 0)
    issues=[]
    if not original_number or original_number == "0": issues.append("رقم البلاغ الأصلي مفقود أو صفري")
    elif numbers[original_number] > 1: issues.append("رقم البلاغ الأصلي مكرر")
    if not original_code: issues.append("رمز البلاغ الأصلي مفقود")
    elif codes[original_code] > 1: issues.append("رمز البلاغ الأصلي مكرر")
    if txt(v[60]).upper() == "NOT OK" or (start is not None and end is not None and end < start): issues.append("قراءة كيلومترات غير سليمة")
    if txt(v[59]).upper() in {"ACTIVE", "NOT OK"}: issues.append("توقيتات البلاغ تحتاج مراجعة")
    if txt(v[61]).upper() == "NOT OK": issues.append("الصف غير مكتمل")
    records.append({
      "incidentNumber":incident_number,"incidentCode":incident_code,"beneficiaryName":txt(v[3]),"gender":txt(v[4]) or "Unknown",
      "contact":txt(v[5]),"age":int(num(v[6])) if num(v[6]) is not None else None,"urgency":txt(v[7]) or "N/A","category":txt(v[8]),"caseType":txt(v[9]) or "Other",
      "pickupType":txt(v[10]),"pickupLocation":txt(v[11]),"pickupGovernorate":txt(v[12]) or "غزة","pickupMunicipality":txt(v[13]),"pickupNeighborhood":txt(v[14]),
      "dropoffType":txt(v[15]),"dropoffLocation":txt(v[16]),"shift":txt(v[17]) or "—","station":txt(v[18]) or "غير محدد","vehicle":txt(v[19]),
      "kmStart":start,"kmEnd":end,"distanceKm":distance,"callDate":day(v[23]),"callTime":clock(v[24]),"dispatchDate":day(v[25]),"dispatchTime":clock(v[26]),
      "onSceneDate":day(v[27]),"arrivalTime":clock(v[28]),"hospitalDate":day(v[29]),"hospitalTime":clock(v[30]),"availableDate":day(v[31]),"clearTime":clock(v[32]),
      "dispatcherPrimary":txt(v[33]),"dispatcherSecondary":txt(v[34]),"emtDriver":txt(v[35]),"emtLead":txt(v[36]),"emtAssist":txt(v[37]),
      "fuelLiters":num(v[38]) or 0,"cancelled":yes(v[39]),"notes":txt(v[40]),"patientDeceased":yes(v[41]),"conflictRelated":yes(v[42]),"chiefComplaint":txt(v[43]),
      "oxygen":yes(v[44]),"bvm":yes(v[45]),"airway":yes(v[46]),"cpr":yes(v[47]),"aed":yes(v[48]),"drugs":yes(v[49]),"woundCare":yes(v[50]),"tourniquet":yes(v[51]),
      "immobilization":yes(v[52]),"cervicalCollar":yes(v[53]),"glucoseCheck":yes(v[54]),"splinting":yes(v[55]),"delivery":yes(v[56]),
      "sourceImportKey":f"datasheet:row:{row}","sourceIncidentNumber":original_number,"sourceIncidentCode":original_code,"importBatch":"Data Sheet Sep 2026",
      "dataQualityStatus":"review" if issues else "ok","dataQualityNotes":"؛ ".join(issues),"approvalStatus":"approved"
    })

target.parent.mkdir(parents=True,exist_ok=True)
target.write_text(json.dumps(records,ensure_ascii=False,separators=(",",":")),encoding="utf-8")
print(json.dumps({"records":len(records),"review":sum(x["dataQualityStatus"]=="review" for x in records),"target":str(target)},ensure_ascii=False))
